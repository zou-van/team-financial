"""Parse monthly Excel files into Markdown and index.json."""

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import yaml
from openpyxl import load_workbook

SHEET_NAME = "差距分析(团队)"
BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_DIR = BASE_DIR / "excel"
DATA_DIR = BASE_DIR / "data"
MAPPING_FILE = BASE_DIR / "team-mapping.yaml"


def extract_month(filepath):
    """Extract YYYY-MM from filename like 'EC团队经营数据汇总202603.xlsx'.

    Returns (month_key, month_label) e.g. ('2026-03', '2026年03月').
    """
    # Find the last 6-digit number in the filename
    matches = re.findall(r"(\d{6})", filepath.name)
    if not matches:
        raise ValueError(f"Cannot find YYYYMM in filename: {filepath.name}")
    yyyymm = matches[-1]
    if len(yyyymm) != 6:
        raise ValueError(f"Invalid date suffix in filename: {filepath.name}")
    yyyy = yyyymm[:4]
    mm = yyyymm[4:6]
    return f"{yyyy}-{mm}", f"{yyyy}年{mm}月"


def load_mapping():
    """Load team-mapping.yaml. Returns None if file missing."""
    if not MAPPING_FILE.exists():
        return None
    with open(MAPPING_FILE, encoding="utf-8") as f:
        return yaml.safe_load(f)


def collect_team_names(mapping):
    """Return set of all expected team names from mapping."""
    names = set()
    if not mapping:
        return names
    names.add(mapping["boss"])
    for sub_teams in mapping.get("leaders", {}).values():
        for sub_team, small_teams in sub_teams.items():
            names.add(sub_team)
            names.update(small_teams)
    return names


def build_column_order(mapping):
    """Return ordered column names following mapping hierarchy."""
    if not mapping:
        return None
    order = []
    for sub_teams in mapping.get("leaders", {}).values():
        for sub_team, small_teams in sub_teams.items():
            order.extend(small_teams)
            order.append(sub_team)
    order.append(mapping["boss"])
    return order


def parse_excel(filepath, mapping):
    """Parse a single Excel file.

    Returns (ordered_headers: list[str], metrics: OrderedDict).
    """
    wb = load_workbook(filepath, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{SHEET_NAME}' not found. Available: {wb.sheetnames}"
        )

    ws = wb[SHEET_NAME]

    # Read column headers from row 1 (cols 2 onward)
    headers = []
    for col in range(2, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        if val is not None:
            headers.append(str(val).strip())

    # Validate against mapping
    expected = collect_team_names(mapping) if mapping else set()
    if expected:
        actual = set(headers)
        missing = expected - actual
        extra = actual - expected
        if missing:
            print(f"  Warning: in mapping but not in Excel: {missing}")
        if extra:
            print(f"  Warning: in Excel but not in mapping: {extra}")

    # Reorder headers by mapping (extra columns appended at end)
    col_order = build_column_order(mapping)
    if col_order:
        header_set = set(headers)
        ordered = [h for h in col_order if h in header_set]
        ordered += [h for h in headers if h not in col_order]
    else:
        ordered = list(headers)

    # Build lookup: header name -> column index (0-based)
    header_to_col = {h: i for i, h in enumerate(headers)}

    # Read data rows
    metrics = OrderedDict()
    for row in range(2, ws.max_row + 1):
        metric_name = ws.cell(row=row, column=1).value
        if metric_name is None:
            continue
        metric_name = str(metric_name).strip()
        if not metric_name:
            continue

        row_values = []
        for h in ordered:
            val = ws.cell(row=row, column=header_to_col[h] + 2).value
            row_values.append(val)
        metrics[metric_name] = row_values

    wb.close()
    return ordered, metrics


def format_value(val):
    """Format a cell value for Markdown display."""
    if val is None:
        return "-"
    if isinstance(val, float):
        if val == int(val):
            return str(int(val))
        return f"{val:.2f}"
    if isinstance(val, int):
        return str(val)
    return str(val)


def generate_markdown(month_label, headers, metrics):
    """Generate Markdown content string."""
    lines = [f"# {month_label} 财务数据\n"]
    for metric_name, values in metrics.items():
        lines.append(f"## {metric_name}")
        lines.append("| " + " | ".join(headers) + " |")
        lines.append("|" + "|".join(["---"] * len(headers)) + "|")
        row = " | ".join(format_value(v) for v in values)
        lines.append(f"| {row} |")
        lines.append("")
    return "\n".join(lines)


def generate_index(months, metrics, mapping):
    """Generate index.json content."""
    return {
        "months": months,
        "metrics": sorted(metrics),
        "teams": mapping or {},
    }


def main():
    DATA_DIR.mkdir(exist_ok=True)
    mapping = load_mapping()
    if mapping is None:
        print("Warning: team-mapping.yaml not found, columns keep Excel order.")

    excel_files = sorted(EXCEL_DIR.glob("*.xlsx"))
    if not excel_files:
        print(f"No .xlsx files found in {EXCEL_DIR}")
        # Still write an empty index.json so frontend works
        index = generate_index([], [], mapping)
        with open(DATA_DIR / "index.json", "w", encoding="utf-8") as f:
            json.dump(index, f, ensure_ascii=False, indent=2)
        print(f"Generated: {DATA_DIR / 'index.json'} (empty)")
        return

    all_metrics = set()
    months_list = []

    for fp in excel_files:
        try:
            month_key, month_label = extract_month(fp)
        except ValueError as e:
            print(f"  ERROR: {e}")
            continue

        print(f"Processing: {fp.name} -> {month_key}")
        try:
            headers, metrics = parse_excel(fp, mapping)
        except Exception as e:
            print(f"  ERROR: {e}")
            continue

        md = generate_markdown(month_label, headers, metrics)
        out_path = DATA_DIR / f"{month_key}.md"
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(md)
        print(
            f"  -> {out_path}  ({len(metrics)} metrics, {len(headers)} teams)"
        )

        months_list.append(month_key)
        all_metrics.update(metrics.keys())

    index = generate_index(sorted(months_list), all_metrics, mapping)
    index_path = DATA_DIR / "index.json"
    with open(index_path, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)
    print(f"Generated: {index_path}")


if __name__ == "__main__":
    main()
